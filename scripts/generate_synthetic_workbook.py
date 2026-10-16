#!/usr/bin/env python3
"""
Generate a SYNTHETIC monthly SG&A workbook for the fictional "APAC Co."

Every number is produced by a seeded random generator. Nothing here is,
or is derived from, any real company's financials. The workbook mirrors
the *shape* of a typical regional FP&A reporting pack (one tab per
function, one column block per market, MTD / YTD / FY x Actual / Plan /
Forecast) so the extraction and commentary pipeline can be demonstrated
end to end.

It also deliberately plants known data-quality defects so the pipeline's
checks have something to catch:
  * two FY 'Delta vs Plan' cells computed against the prior forecast with
    a flipped sign (the sign-error bug the extractor corrects and flags)
  * an 'Unidentified adjustment' line the commentary must flag, not explain

Usage: python scripts/generate_synthetic_workbook.py [out.xlsx] [--month 9]
"""
import argparse
import os
import random
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sga import layout, rules  # noqa: E402

YEAR = 2026
DEPT_ANNUAL_PLAN = {  # USD, fictional
    "OPS": 6_200_000, "TECH": 8_400_000, "PRODUCT": 3_100_000,
    "FINANCE": 2_600_000, "CARE": 5_300_000, "HR": 1_500_000,
    "LEGAL": 1_250_000, "RISK": 1_050_000,
}
GEO_SHARE = {"AU": 0.37, "SG": 0.23, "TW": 0.15, "TH": 0.12, "MY": 0.13}

# (label, account, share of dept plan, outline level)
HC_LINES = [
    ("Salaries & Benefits", "6100", 0.46, 0),
    ("Contract Labour", "6150", 0.08, 0),
    ("Recruiting & Training", "6180", 0.03, 0),
    ("Allocations - Payroll", "6190", 0.04, 1),
]
NONHC_LINES = [
    ("Travel & Entertainment", "6300", 0.05, 0),
    ("Allocations - T&E", "6390", 0.01, 1),
    ("Professional Fees", "6400", 0.08, 0),
    ("Software & Cloud", "6500", 0.12, 0),
    ("Rent & Facilities", "6600", 0.08, 0),
    ("Marketing & Advertising", "6700", 0.03, 0),
    ("Other Opex", "6900", 0.02, 0),
    ("Unidentified adjustment", "6999", 0.0, 0),
]
# (department, geo, line) cells that get the planted sign-error bug
SIGN_ERROR_CELLS = {("HR", "AU", "Total HC Related"),
                    ("LEGAL", "TW", "Salaries & Benefits")}
RAW_FIELDS = [f for f in layout.FIELDS if not f.endswith("delta")]


def leaf_figures(rng, annual_plan, month, volatility, bias=0.0):
    """The raw (non-delta) fields for one line x one market."""
    monthly_plan = annual_plan / 12
    drift = bias + rng.gauss(0, volatility)   # persistent over/under-run
    ytd_plan = monthly_plan * month
    ytd_act = ytd_plan * (1 + drift + rng.gauss(0, volatility / 4))
    fy_fcst = ytd_act + monthly_plan * (12 - month) * (1 + drift * 0.6)
    return {
        "mtd_actual": monthly_plan * (1 + drift + rng.gauss(0, volatility / 2)),
        "mtd_plan": monthly_plan,
        "ytd_actual": ytd_act, "ytd_plan": ytd_plan,
        "fy_prior_forecast": fy_fcst * (1 + rng.gauss(0, 0.015)),
        "fy_forecast": fy_fcst, "fy_plan": annual_plan,
    }


def add_deltas(v):
    v = {k: round(x) for k, x in v.items()}
    v["mtd_delta"] = v["mtd_plan"] - v["mtd_actual"]
    v["ytd_delta"] = v["ytd_plan"] - v["ytd_actual"]
    v["fy_delta"] = v["fy_plan"] - v["fy_forecast"]
    return v


def sum_figs(rows):
    rows = list(rows)
    return {k: sum(r[k] for r in rows) for k in RAW_FIELDS}


def build(out_path, month=9, seed=42):
    rng = random.Random(seed)
    wb = openpyxl.Workbook()
    ctl = wb.active
    ctl.title = "Control"
    for r, (k, v) in enumerate([
            ("Company", "APAC Co. (fictional)"),
            ("Data", "SYNTHETIC - randomly generated, not real financials"),
            ("Year", YEAR), ("Month", month)], start=1):
        ctl.cell(r, 1, k)
        ctl.cell(r, 3, v)

    banner = PatternFill("solid", fgColor="FFF2CC")
    for dkey in rules.DEPARTMENT_ORDER:
        cfg = rules.DEPARTMENTS[dkey]
        ws = wb.create_sheet(cfg["sheet"])
        c = ws.cell(1, 2, f"APAC Co. - {cfg['label']} SG&A - {YEAR}-{month:02d}"
                          "  |  SYNTHETIC DEMO DATA")
        c.font, c.fill = Font(bold=True), banner
        ws.cell(6, layout.COL_LABEL, "Line")
        for geo in rules.GEOS:
            c0 = layout.block_start(geo)
            ws.cell(5, c0, rules.GEO_LABELS[geo]).font = Font(bold=True)
            for f, off in layout.FIELD_OFFSETS.items():
                ws.cell(6, c0 + off, f.replace("_", " ").upper())

        volatility = rng.uniform(0.06, 0.12)
        bias = rng.gauss(0, 0.05)    # department-wide over/under-spend
        state = {"row": layout.DATA_START_ROW}
        sections = {"HC": [], "NONHC": []}

        def write(label, account, by_geo, outline=0, bold=False):
            r = state["row"]
            ws.cell(r, layout.COL_CODE, dkey)
            ws.cell(r, layout.COL_ACCOUNT, account)
            ws.cell(r, layout.COL_LABEL, label).font = Font(bold=bold)
            for geo, v in by_geo.items():
                v = dict(v)
                if (dkey, geo, label) in SIGN_ERROR_CELLS:
                    # planted bug: computed vs prior forecast, sign flipped
                    v["fy_delta"] = v["fy_forecast"] - v["fy_prior_forecast"]
                for f, off in layout.FIELD_OFFSETS.items():
                    ws.cell(r, layout.block_start(geo) + off, v[f])
            if outline:
                ws.row_dimensions[r].outlineLevel = outline
            state["row"] += 1

        def line(label, account, share, outline, section):
            by_geo = {}
            for geo in rules.GEO_LEAVES:
                if label == "Unidentified adjustment":
                    raw = {k: 0 for k in RAW_FIELDS}
                    if geo == "TH" and dkey in ("OPS", "CARE"):
                        hit = rng.randint(18_000, 45_000)
                        raw.update(ytd_actual=hit, fy_forecast=hit,
                                   fy_prior_forecast=hit)
                else:
                    plan = DEPT_ANNUAL_PLAN[dkey] * share * GEO_SHARE[geo]
                    raw = leaf_figures(rng, plan, month, volatility, bias)
                by_geo[geo] = add_deltas(raw)
            by_geo = {"APAC_TOTAL": add_deltas(sum_figs(by_geo.values())),
                      **by_geo}
            write(label, account, by_geo, outline)
            sections[section].append(by_geo)

        def total(label, groups):
            by_geo = {g: add_deltas(sum_figs(grp[g] for grp in groups))
                      for g in rules.GEOS}
            write(label, "", by_geo, bold=True)
            return by_geo

        for spec in HC_LINES:
            line(*spec, "HC")
        hc = total(rules.ROLLUP_LABELS["total_hc"], sections["HC"])
        for spec in NONHC_LINES:
            line(*spec, "NONHC")
        nonhc = total(rules.ROLLUP_LABELS["total_non_hc"], sections["NONHC"])
        total(rules.ROLLUP_LABELS["total_costs"], [hc, nonhc])

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("out", nargs="?", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "..", "data", "synthetic",
        "APAC_Co_SGA_2026-09_SYNTHETIC.xlsx"))
    ap.add_argument("--month", type=int, default=9)
    ap.add_argument("--seed", type=int, default=42)
    a = ap.parse_args()
    print("wrote", build(a.out, a.month, a.seed))
