#!/usr/bin/env python3
"""
Monthly workbook -> SQLite, with cell-level provenance.

Every number is read straight from a cell or computed with one documented
rule (the FY sign-error correction in layout.py). Every stored fact keeps
its sheet and row, so the dashboard can show exactly which cell a figure
came from. Data-quality findings go to `validation_issues`; nothing is
fixed silently.

Usage: python -m sga.extract path/to/workbook.xlsx [db.sqlite]
"""
from __future__ import annotations

import os
import sqlite3
import sys
from datetime import datetime, timezone

import openpyxl

from . import layout, rules

SCHEMA = """
CREATE TABLE IF NOT EXISTS facts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    department TEXT NOT NULL, geo TEXT NOT NULL,
    line TEXT NOT NULL, account TEXT,
    bucket TEXT NOT NULL,              -- HC | NONHC | TOTAL
    outline_level INTEGER DEFAULT 0,   -- 0 = main line, >0 = grouped detail
    mtd_actual REAL, mtd_plan REAL, mtd_delta REAL,
    ytd_actual REAL, ytd_plan REAL, ytd_delta REAL,
    fy_prior_forecast REAL, fy_forecast REAL, fy_plan REAL, fy_delta REAL,
    sign_error_corrected INTEGER DEFAULT 0,
    source_sheet TEXT, source_row INTEGER
);
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS validation_issues (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    department TEXT, geo TEXT, issue TEXT, detail TEXT
);
"""


def _bucket(row, hc_row, nonhc_row):
    if row in (hc_row, nonhc_row):
        return "TOTAL"
    if hc_row and row < hc_row:
        return "HC"
    if nonhc_row and row < nonhc_row:
        return "NONHC"
    return "TOTAL"


def read_control(wb):
    meta = {}
    for label, _, value in wb["Control"].iter_rows(
            min_row=1, max_row=20, max_col=3, values_only=True):
        if label:
            meta[str(label).strip().lower().replace(" ", "_")] = str(value)
    return meta


def _row_values(ws, r):
    return next(ws.iter_rows(min_row=r, max_row=r, max_col=layout.last_col(),
                             values_only=True))


def extract_department(conn, wb, dkey, issues):
    cfg = rules.department_config(dkey)
    ws = wb[cfg["sheet"]]
    labels = {}
    for r in range(layout.DATA_START_ROW, ws.max_row + 1):
        v = ws.cell(r, layout.COL_LABEL).value
        if v:
            labels[str(v).strip()] = r
    hc_row = labels.get(rules.ROLLUP_LABELS["total_hc"])
    nonhc_row = labels.get(rules.ROLLUP_LABELS["total_non_hc"])
    costs_row = labels.get(rules.ROLLUP_LABELS["total_costs"])
    if costs_row is None:
        issues.append((dkey, None, "missing_total_costs_row", cfg["sheet"]))
        return

    cols = ", ".join(layout.FIELDS)
    marks = ",".join("?" * (len(layout.FIELDS) + 9))
    for r in range(layout.DATA_START_ROW, costs_row + 1):
        values = _row_values(ws, r)
        label = values[layout.COL_LABEL - 1]
        if not label:
            continue
        label = str(label).strip()
        dim = ws.row_dimensions.get(r)
        outline = (dim.outlineLevel or 0) if dim else 0
        for g in layout.read_row(values):
            conn.execute(
                f"INSERT INTO facts (department, geo, line, account, bucket, "
                f"outline_level, {cols}, sign_error_corrected, source_sheet, "
                f"source_row) VALUES ({marks})",
                (dkey, g.geo, label, values[layout.COL_ACCOUNT - 1],
                 _bucket(r, hc_row, nonhc_row), outline,
                 *[g.values[f] for f in layout.FIELDS],
                 int(g.sign_error_corrected), cfg["sheet"], r))
            if g.sign_error_corrected:
                issues.append((dkey, g.geo, "sign_error_corrected",
                               f"{cfg['sheet']}!"
                               f"{layout.cell_ref(r, g.geo, 'fy_delta')} "
                               f"({label})"))

    # reconciliation: APAC Total must equal the sum of the markets
    figs = {g.geo: g.values for g in layout.read_row(_row_values(ws, costs_row),
                                                     correct=False)}
    for f in ("ytd_actual", "fy_forecast"):
        leaves = sum(figs[g][f] or 0 for g in rules.GEO_LEAVES)
        total = figs["APAC_TOTAL"][f] or 0
        if abs(leaves - total) > 1:
            issues.append((dkey, "APAC_TOTAL", "total_ne_sum_of_markets",
                           f"Total Costs {f}: total={total:,.0f} vs "
                           f"sum of markets={leaves:,.0f}"))


def run(workbook_path, db_path=":memory:"):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    for t in ("facts", "meta", "validation_issues"):
        conn.execute(f"DELETE FROM {t}")
    meta = {**read_control(wb),
            "source_workbook": os.path.basename(workbook_path),
            "extracted_at": datetime.now(timezone.utc).isoformat()}
    conn.executemany("INSERT INTO meta VALUES (?,?)", meta.items())
    issues = []
    for dkey in rules.DEPARTMENT_ORDER:
        extract_department(conn, wb, dkey, issues)
    conn.executemany("INSERT INTO validation_issues (department, geo, issue, "
                     "detail) VALUES (?,?,?,?)", issues)
    conn.commit()
    return conn


if __name__ == "__main__":
    c = run(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "sga.sqlite")
    n = c.execute("SELECT COUNT(*) FROM facts").fetchone()[0]
    i = c.execute("SELECT COUNT(*) FROM validation_issues").fetchone()[0]
    print(f"{n} facts, {i} validation issues")
